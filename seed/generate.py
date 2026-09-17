#!/usr/bin/env python3
"""
Northwind Materials — synthetic identity source data generator.

Deterministically seeded at 42. Produces the six source exports specified in
Appendix A of the BalkanID FDE/CS hiring task, plus seed/ground_truth.json
recording the correct resolution of every defect injected.

Do NOT hand-clean this output. The mess is the assignment.

Usage:
    python3 seed/generate.py [--out-dir seed/baseline]
"""
from __future__ import annotations

import argparse
import csv
import io
import json
import random
import unicodedata
from datetime import datetime, timedelta, date, timezone
from pathlib import Path

from faker import Faker

SEED = 42
random.seed(SEED)
fake = Faker()
Faker.seed(SEED)

# Fixed "as-of" instant the whole dataset is generated relative to. Every date/datetime
# field that would otherwise be anchored to real wall-clock "now"/"today" must be computed
# relative to AS_OF instead -- otherwise re-running the generator on a different day (or even
# a different second) silently breaks the seed-42 determinism requirement.
AS_OF_DATE = date(2026, 9, 1)
AS_OF_DT = datetime(2026, 9, 1, 9, 0, 0, tzinfo=timezone.utc)


def rel_date(days_ago: int) -> date:
    return AS_OF_DATE - timedelta(days=days_ago)


def rel_dt(days_ago: int) -> datetime:
    return AS_OF_DT - timedelta(days=days_ago)

# ---------------------------------------------------------------------------
# Reference data
# ---------------------------------------------------------------------------

DEPARTMENTS = [
    "Manufacturing", "Engineering", "Sales", "Finance", "Information Technology",
    "Human Resources", "Legal", "Operations", "Quality Assurance",
    "Supply Chain", "Executive", "Customer Success",
]

LOCATIONS = [
    "Detroit, MI, USA", "Cleveland, OH, USA", "Houston, TX, USA",
    "Atlanta, GA, USA", "Pune, India", "Chennai, India",
]

IC_TITLES = {
    "Manufacturing": ["Machine Operator", "Plant Technician", "Production Associate", "Line Supervisor"],
    "Engineering": ["Software Engineer", "Senior Software Engineer", "QA Engineer", "DevOps Engineer"],
    "Sales": ["Account Executive", "Sales Associate", "Sales Engineer"],
    "Finance": ["Financial Analyst", "Accountant", "AP Specialist"],
    "Information Technology": ["IT Support Specialist", "Systems Administrator", "Network Engineer"],
    "Human Resources": ["HR Generalist", "Recruiter", "HR Coordinator"],
    "Legal": ["Paralegal", "Compliance Analyst"],
    "Operations": ["Operations Analyst", "Logistics Coordinator"],
    "Quality Assurance": ["Quality Inspector", "QA Analyst"],
    "Supply Chain": ["Procurement Specialist", "Supply Chain Analyst"],
    "Executive": ["Executive Assistant"],
    "Customer Success": ["Customer Success Associate", "Support Specialist"],
}

TOTAL_WORKERS = 2400
NUM_CONTRACTORS = 180  # AD-only, never in Zoho
TZ_US = timezone(timedelta(hours=-5))   # US Eastern-ish, fixed offset for determinism
TZ_IN = timezone(timedelta(hours=5, minutes=30))  # India Standard Time


def excel_serial(d: date) -> int:
    """Convert a date to an Excel (1900 date system) serial number."""
    epoch = date(1899, 12, 30)  # Excel's fake epoch, accounts for the 1900 leap-year bug
    return (d - epoch).days


def windows_filetime(dt: datetime) -> int:
    """Convert a UTC datetime to a Windows FILETIME integer (100ns ticks since 1601-01-01)."""
    filetime_epoch = datetime(1601, 1, 1, tzinfo=timezone.utc)
    delta = dt.astimezone(timezone.utc) - filetime_epoch
    return int(delta.total_seconds() * 10_000_000)


def rand_date(rng: random.Random, start: date, end: date) -> date:
    days = (end - start).days
    return start + timedelta(days=rng.randint(0, max(days, 0)))


# ---------------------------------------------------------------------------
# Step 1: Build the org chart / worker population
# ---------------------------------------------------------------------------

class Worker:
    __slots__ = (
        "seq", "employee_id", "full_name", "first", "last", "work_email",
        "title", "department", "manager_seq", "location", "hire_date",
        "termination_date", "employment_type", "status", "is_contractor",
        "level",
    )


def build_org_chart(rng: random.Random) -> list[Worker]:
    workers: list[Worker] = []

    def make_worker(seq, title, dept, level, manager_seq, loc=None):
        w = Worker()
        w.seq = seq
        first = fake.first_name()
        last = fake.last_name()
        w.first, w.last = first, last
        w.full_name = f"{first} {last}"
        w.employee_id = f"E{seq:05d}"
        w.work_email = f"{first.lower()}.{last.lower()}@northwindmaterials.com"
        w.title = title
        w.department = dept
        w.manager_seq = manager_seq
        w.location = loc or rng.choice(LOCATIONS)
        w.level = level
        w.is_contractor = False
        workers.append(w)
        return w

    seq = 1
    ceo = make_worker(seq, "Chief Executive Officer", "Executive", 0, None, "Detroit, MI, USA")
    seq += 1

    vps = []
    for dept in DEPARTMENTS:
        if dept == "Executive":
            continue
        v = make_worker(seq, f"VP of {dept}", dept, 1, ceo.seq)
        vps.append(v)
        seq += 1

    directors = []
    n_directors = 30
    for i in range(n_directors):
        vp = rng.choice(vps)
        d = make_worker(seq, f"Director of {vp.department}", vp.department, 2, vp.seq)
        directors.append(d)
        seq += 1

    managers = []
    n_managers = 150
    for i in range(n_managers):
        d = rng.choice(directors)
        m = make_worker(seq, f"Manager, {d.department}", d.department, 3, d.seq)
        managers.append(m)
        seq += 1

    while seq <= TOTAL_WORKERS:
        m = rng.choice(managers)
        title = rng.choice(IC_TITLES[m.department])
        make_worker(seq, title, m.department, 4, m.seq)
        seq += 1

    by_seq = {w.seq: w for w in workers}

    # Hire dates: staggered over the last 12 years, seniority roughly correlated with level
    today = AS_OF_DATE
    for w in workers:
        min_years_ago = {0: 12, 1: 8, 2: 6, 3: 4, 4: 0}[w.level]
        start = today - timedelta(days=365 * 12)
        end = today - timedelta(days=365 * min_years_ago)
        if end < start:
            end = start
        w.hire_date = rand_date(rng, start, end)
        w.termination_date = None
        w.status = "Active"
        w.employment_type = "Full-Time"

    # Terminations: ~8% of the population left over the last 2 years
    n_terminated = int(TOTAL_WORKERS * 0.08)
    terminated = rng.sample(workers, n_terminated)
    for w in terminated:
        w.termination_date = rand_date(rng, today - timedelta(days=730), today - timedelta(days=1))
        w.status = "Terminated"

    return workers, by_seq


# ---------------------------------------------------------------------------
# Step 2: Defect injection helpers for Zoho export
# ---------------------------------------------------------------------------

def format_employee_id_variant(rng: random.Random, canonical: str, variant_pool: dict) -> str:
    """canonical like 'E00042' -> pick a formatting variant, tracked for ground truth."""
    numeric = str(int(canonical[1:]))
    choice = rng.choice(["E0042style", "E42style", "bare", "canonical", "whitespace"])
    if choice == "canonical":
        out = canonical
    elif choice == "E0042style":
        out = "E" + numeric.zfill(4)
    elif choice == "E42style":
        out = "E" + numeric
    elif choice == "bare":
        out = numeric
    else:
        out = canonical + "  "
    variant_pool[out] = canonical
    return out


NON_ASCII_NAMES = [("Priyanka", "Müller"), ("José", "Fernández")]


def build_zoho_rows(workers: list[Worker], rng: random.Random, ground_truth: dict):
    ground_truth["zoho"] = {"employee_id_variants": {}, "status_date_conflicts": [],
                             "blank_manager": [], "terminated_manager": [],
                             "id_collisions": [], "non_ascii": [], "unquoted_comma_row": None}

    rows = []
    id_variant_pool = ground_truth["zoho"]["employee_id_variants"]

    # 120 rows: termination_date in the past but status still Active (data defect)
    terminated_workers = [w for w in workers if w.termination_date is not None]
    conflict_sample = rng.sample(terminated_workers, min(120, len(terminated_workers)))
    conflict_seqs = {w.seq for w in conflict_sample}

    # 60 blank manager_employee_id
    non_ceo = [w for w in workers if w.manager_seq is not None]
    blank_mgr_sample = set(w.seq for w in rng.sample(non_ceo, 60))

    # 25 more whose manager is themselves terminated -- force by reassigning manager
    terminated_seqs = {w.seq for w in terminated_workers}
    candidates_for_term_mgr = [w for w in non_ceo if w.seq not in blank_mgr_sample][:500]
    term_mgr_sample = rng.sample(candidates_for_term_mgr, min(25, len(candidates_for_term_mgr)))
    term_mgr_seqs = set()
    for w in term_mgr_sample:
        if terminated_workers:
            fake_mgr = rng.choice(terminated_workers)
            w.manager_seq = fake_mgr.seq
            term_mgr_seqs.add(w.seq)

    # 3 employee_id collisions: pick 3 pairs of distinct people sharing one employee_id
    collision_targets = rng.sample(workers, 6)
    collision_map = {}  # shared_id -> [seqs]
    for i in range(0, 6, 2):
        shared_id = collision_targets[i].employee_id
        collision_map[shared_id] = [collision_targets[i].seq, collision_targets[i + 1].seq]
    ground_truth["zoho"]["id_collisions"] = [
        {"shared_employee_id": k, "true_worker_seqs": v} for k, v in collision_map.items()
    ]
    collided_second_seqs = {pair[1] for pair in collision_map.values()}

    # 2 non-ASCII names + 1 unquoted comma
    non_ascii_sample = rng.sample(workers, 2)
    for w, (f, l) in zip(non_ascii_sample, NON_ASCII_NAMES):
        w.first, w.last = f, l
        w.full_name = f"{f} {l}"
        # Recompute work_email to match the new name -- it was already set once in
        # build_org_chart from the OLD (pre-override) name, and every other source (AD, Entra,
        # Salesforce) derives its own identifiers from w.first/w.last directly. Leaving the
        # stale email in place would silently break every downstream correlation for these two
        # people (a name/email mismatch nobody intended), instead of testing the intended
        # defect (a non-ASCII name that should otherwise resolve normally).
        w.work_email = f"{f.lower()}.{l.lower()}@northwindmaterials.com"
    ground_truth["zoho"]["non_ascii"] = [w.employee_id for w in non_ascii_sample]

    comma_worker = rng.choice([w for w in workers if w not in non_ascii_sample])
    comma_worker.full_name = f"{comma_worker.last}, Jr. {comma_worker.first}"  # unquoted comma injected at write time
    ground_truth["zoho"]["unquoted_comma_row"] = comma_worker.employee_id

    for w in workers:
        emp_id_out = format_employee_id_variant(rng, w.employee_id, id_variant_pool)
        if w.seq in collided_second_seqs:
            for shared_id, pair in collision_map.items():
                if w.seq == pair[1]:
                    emp_id_out = shared_id

        manager_out = ""
        if w.seq in blank_mgr_sample:
            manager_out = ""
            ground_truth["zoho"]["blank_manager"].append(w.employee_id)
        elif w.manager_seq is not None:
            mgr = next(x for x in workers if x.seq == w.manager_seq)
            manager_out = mgr.employee_id
            if w.seq in term_mgr_seqs:
                ground_truth["zoho"]["terminated_manager"].append(w.employee_id)

        status_out = w.status
        if w.seq in conflict_seqs:
            status_out = "Active"  # conflict: termination_date set but status left Active
            ground_truth["zoho"]["status_date_conflicts"].append(w.employee_id)

        # date format rotation: ISO / US-slash / Excel serial
        fmt = rng.choice(["iso", "us", "excel"])
        if fmt == "iso":
            hire_out = w.hire_date.isoformat()
        elif fmt == "us":
            hire_out = w.hire_date.strftime("%m/%d/%Y")
        else:
            hire_out = str(excel_serial(w.hire_date))

        term_out = ""
        if w.termination_date:
            term_out = w.termination_date.isoformat()

        rows.append({
            "employee_id": emp_id_out,
            "full_name": w.full_name,
            "work_email": w.work_email,
            "title": w.title,
            "department": w.department,
            "manager_employee_id": manager_out,
            "location": w.location,
            "hire_date": hire_out,
            "termination_date": term_out,
            "employment_type": w.employment_type,
            "status": status_out,
            "_true_employee_id": w.employee_id,  # not written to CSV; used only for ground truth cross-linking
        })

    return rows


def write_zoho_csv(rows: list[dict], path: Path):
    fieldnames = ["employee_id", "full_name", "work_email", "title", "department",
                  "manager_employee_id", "location", "hire_date", "termination_date",
                  "employment_type", "status"]
    lines = [",".join(fieldnames)]
    for r in rows:
        vals = []
        for f in fieldnames:
            v = str(r[f])
            if f == "full_name" and "," in v:
                # inject deliberately WITHOUT quoting -- this is the malformed-CSV defect
                vals.append(v)
            elif "," in v:
                vals.append(f'"{v}"')
            else:
                vals.append(v)
        lines.append(",".join(vals))
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


# ---------------------------------------------------------------------------
# Step 3: Entra ID export
# ---------------------------------------------------------------------------

def build_entra(workers: list[Worker], rng: random.Random, ground_truth: dict):
    ground_truth["entra"] = {"null_employee_id": [], "display_name_mismatch": {}, "dual_account_pairs": []}
    active_pool = [w for w in workers]
    n = 2150
    chosen = rng.sample(active_pool, min(n, len(active_pool)))

    null_emp_id_sample = set(w.seq for w in rng.sample(chosen, int(len(chosen) * 0.15)))
    mismatch_sample = rng.sample(chosen, 40)

    users = []
    for w in chosen:
        upn = f"{w.first.lower()}.{w.last.lower()}@northwindmaterials.onmicrosoft.com"
        display = w.full_name
        if w in mismatch_sample:
            display = f"{rng.choice(['Ms.', 'Mr.'])} {w.last}-{fake.last_name()}"  # married-name style mismatch
            ground_truth["entra"]["display_name_mismatch"][w.employee_id] = display

        last_signin = None
        if w.status == "Active" and rng.random() > 0.05:
            last_signin = fake.date_time_between(start_date=rel_dt(60), end_date=AS_OF_DT).isoformat() + "Z"

        emp_id_field = None if w.seq in null_emp_id_sample else w.employee_id
        if w.seq in null_emp_id_sample:
            ground_truth["entra"]["null_employee_id"].append(w.full_name)

        users.append({
            "id": f"entra-{w.seq:06d}",
            "userPrincipalName": upn,
            "displayName": display,
            "mail": w.work_email,
            "accountEnabled": w.status == "Active",
            "employeeId": emp_id_field,
            "lastSignInDateTime": last_signin,
            "assignedGroups": [],
            "appRoleAssignments": [],
            "_true_employee_id": w.employee_id,
        })

    # 12 people hold both Entra and AD accounts under different UPNs -- pick 12 from chosen
    dual_sample = rng.sample(users, 12)
    for u in dual_sample:
        ground_truth["entra"]["dual_account_pairs"].append(u["_true_employee_id"])

    return users, {u["_true_employee_id"] for u in dual_sample}


# ---------------------------------------------------------------------------
# Step 4: Active Directory export
# ---------------------------------------------------------------------------

SERVICE_PREFIXES = ["svc-", "plant-shared-"]

def build_ad(workers: list[Worker], rng: random.Random, dual_account_seqs: set, ground_truth: dict):
    ground_truth["ad"] = {"contractors": [], "service_accounts": [], "orphans": [],
                           "nested_groups": {}, "dual_account_upns": {}, "shared_accounts": []}

    rows = []

    # Employees who also appear in AD (most of them), using standard UPN except the 12 dual-account people.
    # The dual-account people MUST be present in this AD sample, or the "12 people hold both an
    # Entra and an AD account" invariant silently breaks.
    dual_by_id = {w.employee_id: w for w in workers if w.employee_id in dual_account_seqs}
    forced = list(dual_by_id.values())
    remaining_pool = [w for w in workers if w.employee_id not in dual_account_seqs]
    emp_sample = forced + rng.sample(remaining_pool, 2090 - len(forced))
    for w in emp_sample:
        if w.employee_id in dual_account_seqs:
            upn = f"{w.first.lower()}{w.last[0].lower()}@northwindmaterials.com"  # deliberately different UPN
            ground_truth["ad"]["dual_account_upns"][w.employee_id] = upn
        else:
            upn = f"{w.first.lower()}.{w.last.lower()}@northwindmaterials.com"

        groups = ["All-Employees"]
        dept_group = f"{w.department.replace(' ', '-')}-All"
        groups.append(dept_group)
        if w.department == "Manufacturing" and rng.random() < 0.3:
            groups.append("Plant-Ops-Leads")  # nested 3 levels deep, resolved by access_graph later
            ground_truth["ad"]["nested_groups"].setdefault(w.employee_id, []).append(
                "Plant-Ops-Leads -> Manufacturing-All -> All-Employees"
            )

        last_logon_dt = fake.date_time_between(start_date=rel_dt(90), end_date=AS_OF_DT)
        rows.append({
            "samAccountName": f"{w.first.lower()}.{w.last.lower()}",
            "distinguishedName": f"CN={w.full_name},OU={w.department.replace(' ', '')},DC=northwind,DC=local",
            "displayName": w.full_name,
            "userPrincipalName": upn,
            "enabled": "TRUE" if w.status == "Active" else "FALSE",
            "lastLogonTimestamp": str(windows_filetime(last_logon_dt)),
            "memberOf": ";".join(groups),
            "description": f"{w.title} - {w.department}",
            "_true_employee_id": w.employee_id,
            "_classification": "employee",
        })

    # 180 contractors, AD-only, never in Zoho
    for i in range(NUM_CONTRACTORS):
        first, last = fake.first_name(), fake.last_name()
        name = f"{first} {last}"
        sam = f"ctr.{first.lower()}.{last.lower()}"
        rows.append({
            "samAccountName": sam,
            "distinguishedName": f"CN={name},OU=Contractors,DC=northwind,DC=local",
            "displayName": name,
            "userPrincipalName": f"{first.lower()}.{last.lower()}@contractor.northwindmaterials.com",
            "enabled": "TRUE",
            "lastLogonTimestamp": str(windows_filetime(fake.date_time_between(start_date=rel_dt(90), end_date=AS_OF_DT))),
            "memberOf": "Contractors-All;All-Employees",
            "description": f"Contractor - {rng.choice(['Staffing Partner A', 'Staffing Partner B'])} - vendor engagement",
            "_true_employee_id": None,
            "_classification": "contractor",
        })
        ground_truth["ad"]["contractors"].append(sam)

    # remaining ~40 unmatched: service accounts + genuine orphans (total unmatched ~220)
    n_service = 28
    n_orphan = 12
    for i in range(n_service):
        prefix = rng.choice(SERVICE_PREFIXES)
        name = f"{prefix}{fake.word()}{i:02d}"
        is_shared = prefix in SERVICE_PREFIXES and i < 6
        rows.append({
            "samAccountName": name,
            "distinguishedName": f"CN={name},OU=ServiceAccounts,DC=northwind,DC=local",
            "displayName": name,
            "userPrincipalName": f"{name}@northwindmaterials.com",
            "enabled": "TRUE",
            "lastLogonTimestamp": str(windows_filetime(fake.date_time_between(start_date=rel_dt(30), end_date=AS_OF_DT))),
            "memberOf": "Service-Accounts",
            "description": "Non-interactive service/CI account" if i % 3 else "Owned by IT Automation team",
            "_true_employee_id": None,
            "_classification": "service",
        })
        ground_truth["ad"]["service_accounts"].append(name)
        if i < 6:
            ground_truth["ad"]["shared_accounts"].append(name)

    # one integration account owned by an employee who left in March
    march_leaver = next((w for w in workers if w.termination_date and w.termination_date.month == 3), workers[0])
    rows.append({
        "samAccountName": "svc-sfdc-integration",
        "distinguishedName": "CN=svc-sfdc-integration,OU=ServiceAccounts,DC=northwind,DC=local",
        "displayName": "svc-sfdc-integration",
        "userPrincipalName": "svc-sfdc-integration@northwindmaterials.com",
        "enabled": "TRUE",
        "lastLogonTimestamp": str(windows_filetime(fake.date_time_between(start_date=rel_dt(10), end_date=AS_OF_DT))),
        "memberOf": "Service-Accounts;Salesforce-Integration",
        "description": f"Integration account, originally set up by {march_leaver.full_name} (terminated)",
        "_true_employee_id": None,
        "_classification": "service_orphaned_owner",
    })
    ground_truth["ad"]["service_accounts"].append("svc-sfdc-integration")
    ground_truth["ad"]["orphaned_owner_leaver"] = march_leaver.employee_id

    for i in range(n_orphan):
        first, last = fake.first_name(), fake.last_name()
        name = f"{first.lower()}.{last.lower()}{i}"
        rows.append({
            "samAccountName": name,
            "distinguishedName": f"CN={first} {last},OU=Users,DC=northwind,DC=local",
            "displayName": f"{first} {last}",
            "userPrincipalName": f"{name}@northwindmaterials.com",
            "enabled": "TRUE" if i % 2 == 0 else "FALSE",
            "lastLogonTimestamp": str(windows_filetime(fake.date_time_between(start_date=rel_dt(400), end_date=rel_dt(200)))),
            "memberOf": "All-Employees",
            "description": "",  # no contractor/service evidence -- genuine orphan
            "_true_employee_id": None,
            "_classification": "orphan",
        })
        ground_truth["ad"]["orphans"].append(name)

    rng.shuffle(rows)
    return rows


def write_ad_csv(rows, path: Path):
    fieldnames = ["samAccountName", "distinguishedName", "displayName", "userPrincipalName",
                  "enabled", "lastLogonTimestamp", "memberOf", "description"]
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fieldnames)
        w.writeheader()
        for r in rows:
            w.writerow({k: r[k] for k in fieldnames})


# ---------------------------------------------------------------------------
# Step 5: Salesforce export
# ---------------------------------------------------------------------------

def build_salesforce(workers: list[Worker], rng: random.Random, ground_truth: dict):
    ground_truth["salesforce"] = {"legacy_domain": [], "active_but_terminated": []}
    active_pop = [w for w in workers if w.department in ("Sales", "Customer Success", "Executive", "Finance")]
    extra = rng.sample([w for w in workers if w not in active_pop], 200)
    pool_by_seq = {w.seq: w for w in (active_pop + extra)}  # dedupe: a worker must appear at most once
    pool = list(pool_by_seq.values())
    chosen = rng.sample(pool, min(480, len(pool)))

    legacy_sample = rng.sample(chosen, 30)
    terminated_pool = [w for w in chosen if w.status == "Terminated"]
    term_active_sample = rng.sample(terminated_pool, min(8, len(terminated_pool)))

    rows = []
    for w in chosen:
        email = w.work_email
        if w in legacy_sample:
            email = f"{w.first.lower()}.{w.last.lower()}@acquiredbrand.com"
            ground_truth["salesforce"]["legacy_domain"].append(w.employee_id)

        is_active = "TRUE"
        if w in term_active_sample:
            is_active = "TRUE"  # SFDC record never deactivated after termination
            ground_truth["salesforce"]["active_but_terminated"].append(w.employee_id)
        elif w.status == "Terminated":
            is_active = "FALSE"

        profile = rng.choice(["Standard User", "Sales Admin", "System Administrator", "Read Only"])
        perm_sets = rng.sample(["Revenue_Admin", "Opportunity_Approver", "Report_Builder", "Data_Export"], k=rng.randint(0, 2))

        rows.append({
            "username": email,
            "name": w.full_name,
            "email": email,
            "profile": profile,
            "permission_sets": "|".join(perm_sets),
            "is_active": is_active,
            "last_login": fake.date_time_between(start_date=rel_dt(120), end_date=AS_OF_DT).isoformat(),
            "_true_employee_id": w.employee_id,
        })
    return rows


def write_salesforce_csv(rows, path: Path):
    fieldnames = ["username", "name", "email", "profile", "permission_sets", "is_active", "last_login"]
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fieldnames)
        w.writeheader()
        for r in rows:
            w.writerow({k: r[k] for k in fieldnames})


# ---------------------------------------------------------------------------
# Step 6: AWS IAM export
# ---------------------------------------------------------------------------

def build_aws(workers: list[Worker], rng: random.Random, ground_truth: dict):
    ground_truth["aws"] = {"no_human_match": [], "stale_access_keys": [], "admin_by_transitivity_role": None}
    accounts = ["111111111111-prod", "222222222222-staging", "333333333333-shared-services"]
    eng_pop = [w for w in workers if w.department in ("Engineering", "Information Technology", "Operations")]
    pool = rng.sample(eng_pop + rng.sample(workers, 60), min(178, len(eng_pop) + 60))

    users = []
    for w in pool:
        key_age_days = rng.randint(1, 500)
        users.append({
            "username": f"{w.first.lower()}.{w.last.lower()}",
            "account": rng.choice(accounts),
            "arn": f"arn:aws:iam::{rng.choice(accounts).split('-')[0]}:user/{w.first.lower()}.{w.last.lower()}",
            "groups": rng.sample(["Developers", "ReadOnly", "Admins", "DataTeam"], k=rng.randint(1, 2)),
            "access_keys": [{"created_days_ago": key_age_days, "active": True}],
            "_true_employee_id": w.employee_id,
        })

    for i in range(22):
        name = f"automation-user-{i:03d}"
        users.append({
            "username": name,
            "account": rng.choice(accounts),
            "arn": f"arn:aws:iam::{rng.choice(accounts).split('-')[0]}:user/{name}",
            "groups": ["CI-CD"],
            "access_keys": [{"created_days_ago": rng.randint(1, 600), "active": True}],
            "_true_employee_id": None,
        })
        ground_truth["aws"]["no_human_match"].append(name)

    stale_sample = rng.sample(users, 4)
    for u in stale_sample:
        u["access_keys"][0]["created_days_ago"] = rng.randint(401, 700)
        ground_truth["aws"]["stale_access_keys"].append(u["username"])

    roles = []
    for i in range(90):
        roles.append({
            "role_name": f"role-{fake.word()}-{i:03d}",
            "account": rng.choice(accounts),
            "trust_policy": "internal",
            "attached_policies": rng.sample(["ReadOnlyAccess", "PowerUserAccess", "AdministratorAccess", "S3FullAccess"], k=1),
        })

    # One transitive admin chain: RoleA assumes RoleB assumes RoleC which has AdministratorAccess
    roles[0]["role_name"] = "role-deploy-automation-000"
    roles[0]["attached_policies"] = ["S3FullAccess"]
    roles[0]["assume_role_chain"] = ["role-deploy-automation-000", "role-ci-bridge-041", "role-break-glass-admin-089"]
    roles[41]["role_name"] = "role-ci-bridge-041"
    roles[41]["attached_policies"] = ["ReadOnlyAccess"]
    roles[89]["role_name"] = "role-break-glass-admin-089"
    roles[89]["attached_policies"] = ["AdministratorAccess"]
    ground_truth["aws"]["admin_by_transitivity_role"] = "role-deploy-automation-000"

    return {"accounts": accounts, "users": users, "roles": roles}


# ---------------------------------------------------------------------------
# Step 7: PlantOps spreadsheet
# ---------------------------------------------------------------------------

def name_variant(rng: random.Random, first: str, last: str) -> str:
    style = rng.choice(["last_first", "first_last_initial", "email_like"])
    if style == "last_first":
        return f"{last}, {first}"
    elif style == "first_last_initial":
        return f"{first} {last[0]}."
    else:
        return f"{first.lower()}.{last.lower()}@"


def build_plantops(workers: list[Worker], rng: random.Random, ground_truth: dict):
    ground_truth["plantops"] = {"name_key_variants": {}, "stale_sheet_note": "Sheet2 is prior-year extract, unlabeled"}
    manu_pop = [w for w in workers if w.department == "Manufacturing"] or workers
    non_manu = [w for w in workers if w.department != "Manufacturing"]
    n_manu_pick = min(len(manu_pop), 210)
    n_other_pick = 310 - n_manu_pick
    pool = rng.sample(manu_pop, n_manu_pick) + rng.sample(non_manu, min(n_other_pick, len(non_manu)))
    rng.shuffle(pool)

    rows = []
    for w in pool:
        variant = name_variant(rng, w.first, w.last)
        ground_truth["plantops"]["name_key_variants"].setdefault(w.employee_id, []).append(variant)
        rows.append({
            "Employee Name": variant,
            "Role": rng.choice(["Plant-Ops-Leads", "Line Access", "Warehouse Access", "Safety Officer"]),
            "Granted By": rng.choice(["S. Ortiz", "Plant Manager", "R. Alvarez"]),
            "Date": fake.date_between(start_date=rel_date(500), end_date=AS_OF_DATE),
        })
    return rows


def write_plantops_xlsx(rows, path: Path):
    from openpyxl import Workbook
    wb = Workbook()
    # openpyxl stamps workbook.properties.created/modified with real wall-clock time by
    # default, which would silently break byte-for-byte determinism across runs.
    wb.properties.created = AS_OF_DT.replace(tzinfo=None)
    wb.properties.modified = AS_OF_DT.replace(tzinfo=None)
    ws = wb.active
    ws.title = "September 2026"

    # Header begins on row 4 (rows 1-3 are title/merged banner), with a merged cell
    ws.merge_cells("A1:D1")
    ws["A1"] = "Northwind Materials -- PlantOps Access Extract"
    ws["A2"] = "Confidential -- internal use only"
    ws["A3"] = "Generated by S. Ortiz"
    headers = ["Employee Name", "Role", "Granted By", "Date"]
    for col, h in enumerate(headers, start=1):
        ws.cell(row=4, column=col, value=h)
    for i, r in enumerate(rows, start=5):
        ws.cell(row=i, column=1, value=r["Employee Name"])
        ws.cell(row=i, column=2, value=r["Role"])
        ws.cell(row=i, column=3, value=r["Granted By"])
        ws.cell(row=i, column=4, value=r["Date"].isoformat())

    # Second, unlabeled worksheet: last year's stale extract
    ws2 = wb.create_sheet("Sheet2")
    ws2["A1"] = "Employee Name"
    ws2["B1"] = "Role"
    ws2["C1"] = "Granted By"
    ws2["D1"] = "Date"
    stale_sample = rows[:40]
    for i, r in enumerate(stale_sample, start=2):
        ws2.cell(row=i, column=1, value=r["Employee Name"])
        ws2.cell(row=i, column=2, value=r["Role"])
        ws2.cell(row=i, column=3, value="Prior Year Extract")
        ws2.cell(row=i, column=4, value="2025-09-01")

    wb.save(path)
    _freeze_xlsx_metadata(path)


def _freeze_xlsx_metadata(path: Path):
    """openpyxl forcibly re-stamps docProps/core.xml's <dcterms:modified> and every zip
    member's timestamp with real wall-clock time inside Workbook.save(), regardless of what
    Workbook.properties.modified was set to beforehand. Rewrite the archive after the fact so
    the file is byte-identical across runs, which the seed-42 determinism requirement demands.
    """
    import zipfile
    import re as _re

    fixed_ts = (2026, 9, 1, 9, 0, 0)
    buf = io.BytesIO()
    with zipfile.ZipFile(path, "r") as zin:
        names = zin.namelist()
        contents = {n: zin.read(n) for n in names}

    core = contents.get("docProps/core.xml")
    if core is not None:
        core = _re.sub(
            rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
            rb"\g<1>2026-09-01T09:00:00Z\g<2>",
            core,
        )
        contents["docProps/core.xml"] = core

    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in names:
            info = zipfile.ZipInfo(n, date_time=fixed_ts)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            zout.writestr(info, contents[n])

    path.write_bytes(buf.getvalue())


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--out-dir", default="seed/baseline")
    args = parser.parse_args()

    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    rng = random.Random(SEED)
    ground_truth: dict = {"seed": SEED, "as_of": AS_OF_DT.isoformat(), "note": "generated_at is intentionally omitted so ground_truth.json is byte-identical across runs"}

    workers, by_seq = build_org_chart(rng)

    zoho_rows = build_zoho_rows(workers, rng, ground_truth)
    write_zoho_csv(zoho_rows, out_dir / "zoho_people_workers.csv")

    entra_users, dual_account_seqs = build_entra(workers, rng, ground_truth)
    (out_dir / "entra_users.json").write_text(
        json.dumps([{k: v for k, v in u.items() if not k.startswith("_")} for u in entra_users], indent=2),
        encoding="utf-8",
    )

    ad_rows = build_ad(workers, rng, dual_account_seqs, ground_truth)
    write_ad_csv(ad_rows, out_dir / "ad_accounts.csv")

    sf_rows = build_salesforce(workers, rng, ground_truth)
    write_salesforce_csv(sf_rows, out_dir / "salesforce_users.csv")

    aws_data = build_aws(workers, rng, ground_truth)
    (out_dir / "aws_iam.json").write_text(json.dumps(aws_data, indent=2), encoding="utf-8")

    plantops_rows = build_plantops(workers, rng, ground_truth)
    write_plantops_xlsx(plantops_rows, out_dir / "plantops_access_202609.xlsx")

    # Provenance ground truth for every zoho row -> canonical employee id (handles formatting variants)
    ground_truth["zoho"]["canonical_employee_ids"] = [w.employee_id for w in workers]
    ground_truth["zoho"]["total_rows"] = len(zoho_rows)
    ground_truth["ad"]["total_rows"] = len(ad_rows)
    ground_truth["entra"]["total_rows"] = len(entra_users)
    ground_truth["salesforce"]["total_rows"] = len(sf_rows)
    ground_truth["plantops"]["total_rows"] = len(plantops_rows)
    ground_truth["aws"]["total_users"] = len(aws_data["users"])
    ground_truth["aws"]["total_roles"] = len(aws_data["roles"])

    (Path("seed") / "ground_truth.json").write_text(json.dumps(ground_truth, indent=2, default=str), encoding="utf-8")

    print(f"Generated {len(zoho_rows)} Zoho rows, {len(entra_users)} Entra users, "
          f"{len(ad_rows)} AD rows, {len(sf_rows)} Salesforce rows, "
          f"{len(aws_data['users'])} AWS IAM users / {len(aws_data['roles'])} roles, "
          f"{len(plantops_rows)} PlantOps rows -> {out_dir}/")
    print("ground_truth.json written to seed/ground_truth.json")


if __name__ == "__main__":
    main()
